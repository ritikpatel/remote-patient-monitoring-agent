"""Export the exact dataset the risk engine's promoted model trains on, as Excel.

One workbook, Train and Test on their own sheets, split by the *same* holdout
`eval/prediction.py` uses -- reused rather than reinvented, so the file
corresponds to a split this project already reports against. Column set is
asserted against `ml/models/promoted/feature_manifest.json` at export time, so
the workbook cannot silently drift from the model actually being served.

Usage:
    python ml/evaluation/export_training_data.py [--out exports/]
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(REPO_ROOT))

from eval.prediction import HOLDOUT_RANDOM_STATE, HOLDOUT_SPLITS, _one_holdout_split  # noqa: E402

from ml.evaluation.run_all import WAREHOUSE_DB  # noqa: E402
from ml.features import engineer, labels  # noqa: E402

MANIFEST = REPO_ROOT / "ml" / "models" / "promoted" / "feature_manifest.json"
HORIZON = 6
ARIAL = "Arial"

NOTICE = (
    "This platform is validated on a 100-patient demo subset of MIMIC-IV. The engineering "
    "is real and the methodology is rigorous; the clinical performance figures demonstrate "
    "pipeline validity and DO NOT transfer to clinical practice (PROJECT_PLAN.md s17)."
)


def build(conn):
    grid = conn.execute("select stay_id, hour from capstone.hourly_grid").fetchdf()
    lab = labels.build_labels(conn, grid, horizons=(HORIZON,))
    features = engineer.build_feature_frame(conn)
    x, y, groups = engineer.feature_matrix_for_training(features, lab, f"label_{HORIZON}h")
    keys = lab[["stay_id", "hour"]].merge(features[["stay_id", "hour"]], how="inner")
    assert len(keys) == len(x), (len(keys), len(x))
    return x, y, groups, keys


def style(ws) -> None:
    fill, font = PatternFill("solid", fgColor="1F3864"), Font(
        name=ARIAL, bold=True, color="FFFFFF", size=10
    )
    for c in ws[1]:
        c.fill, c.font = fill, font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(
            max(len(str(col[0].value or "")) + 3, 10), 40
        )


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path, default=REPO_ROOT / "exports")
    ap.add_argument("--db", type=Path, default=WAREHOUSE_DB)
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    conn = duckdb.connect(str(args.db), read_only=True)
    x, y, groups, keys = build(conn)
    conn.close()

    manifest = json.loads(MANIFEST.read_text())
    assert manifest["horizon_h"] == HORIZON, manifest["horizon_h"]
    assert list(x.columns) == manifest["feature_columns"], (
        "exported columns do not match the served model:\n"
        f"  only in export:   {sorted(set(x.columns) - set(manifest['feature_columns']))}\n"
        f"  only in manifest: {sorted(set(manifest['feature_columns']) - set(x.columns))}"
    )
    print(f"OK: {len(x.columns)} columns match promoted model '{manifest['model_name']}'")

    train_idx, test_idx = _one_holdout_split(y, groups)
    out = x.copy()
    out.insert(0, f"label_{HORIZON}h", y.to_numpy())
    out.insert(0, "subject_id", groups.to_numpy())
    out.insert(0, "hour", keys["hour"].to_numpy())
    out.insert(0, "stay_id", keys["stay_id"].to_numpy())
    train, test = out.iloc[train_idx].copy(), out.iloc[test_idx].copy()
    assert not (set(train.subject_id) & set(test.subject_id)), "subject leak across the split"
    print("OK: no subject appears in both Train and Test")

    label = f"label_{HORIZON}h"
    feat_cols = [c for c in out.columns if c not in ("stay_id", "hour", "subject_id", label)]
    fdict = pd.DataFrame(
        {
            "feature": feat_cols,
            "dtype": [str(train[c].dtype) for c in feat_cols],
            "non_null_pct_train": [round(train[c].notna().mean() * 100, 1) for c in feat_cols],
        }
    )

    path = args.out / "risk_engine_training_data.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame().to_excel(xw, sheet_name="README", index=False)
        train.to_excel(xw, sheet_name="Train", index=False)
        test.to_excel(xw, sheet_name="Test", index=False)
        fdict.to_excel(xw, sheet_name="Feature_Dictionary", index=False)

    wb = load_workbook(path)
    ws = wb["README"]
    ws.delete_rows(1)
    ws["A1"] = "Risk engine — training and test data"
    ws["A1"].font = Font(name=ARIAL, bold=True, size=14)
    ws["A2"] = NOTICE
    ws["A2"].font = Font(name=ARIAL, size=9, italic=True)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 46
    for c in ("A2", "B2", "C2", "D2"):
        ws[c].fill = PatternFill("solid", fgColor="FFF2CC")

    rows = [
        ("h", "What this is", ""),
        (
            "t",
            "Model served",
            f"{manifest['model_name']} - the promoted model that services/risk-engine "
            "serves at /score/ml.",
        ),
        (
            "t",
            "Prediction task",
            "Composite deterioration (death, vasopressor initiation, invasive ventilation "
            f"initiation, or unplanned ICU readmission) within the next {HORIZON}h.",
        ),
        (
            "t",
            "One row =",
            "one patient-hour from capstone.hourly_grid, censored to rows strictly BEFORE "
            "the stay's first composite event.",
        ),
        (
            "t",
            "Features",
            f"{len(x.columns)} columns, asserted at export time to equal "
            "ml/models/promoted/feature_manifest.json.",
        ),
        ("h", "The feature set was pruned", ""),
        (
            "t",
            "90 -> 67 columns",
            "news2, sofa_24hours, the 18 rolling means and the 3 lab-ordering-intensity "
            "columns were removed after measurement. AUPRC rose 0.398 -> 0.493, winning 17 "
            "of 20 paired CV repeats. See ml/evaluation/feature_pruning_report.md.",
        ),
        (
            "t",
            "Why news2/sofa left",
            "They are deterministic functions of vitals already present (measured "
            "contribution 0.000), and SOFA's cardiovascular component is scored on "
            "vasopressor dose - which is one of the labels. They remain the BASELINES the "
            "model must beat, scored from their own matrix.",
        ),
        (
            "t",
            "Read this as",
            "a smaller model that is better mostly because 90 features against 120 "
            "positives was over-parameterised: variance removed, not signal added.",
        ),
        ("h", "How the split was made", ""),
        (
            "t",
            "Method",
            f"StratifiedGroupKFold(n_splits={HOLDOUT_SPLITS}, shuffle=True, "
            f"random_state={HOLDOUT_RANDOM_STATE}), first fold held out - "
            "eval/prediction.py's own holdout, reused rather than reinvented.",
        ),
        (
            "t",
            "Grouped by",
            "subject_id, NOT stay_id. 21 of 93 subjects have >1 ICU stay; splitting by "
            "stay let one patient sit in train and test at once (finding F6, worth "
            "+0.0499 AUPRC of false optimism).",
        ),
        ("t", "Leak check", "Verified at export: no subject_id appears in both sheets."),
        ("h", "Caveats", ""),
        (
            "t",
            "One split, not the protocol",
            "Headline figures come from 20-repeat grouped CV (ml/evaluation/report.md). "
            "Metrics computed on this single split will differ.",
        ),
        (
            "t",
            "Confidence intervals",
            "With 49 positive subjects the AUPRC 95% CI is ~0.25 wide and does not close. "
            "See ml/evaluation/reliability_report.md.",
        ),
        ("h", "Counts (computed at export)", ""),
        ("n", "Train rows", f"{len(train)}"),
        ("n", "Train positives", f"{int(train[label].sum())}  ({train[label].mean():.1%})"),
        ("n", "Train subjects", f"{train.subject_id.nunique()}"),
        ("n", "Test rows", f"{len(test)}"),
        ("n", "Test positives", f"{int(test[label].sum())}  ({test[label].mean():.1%})"),
        ("n", "Test subjects", f"{test.subject_id.nunique()}"),
        ("h", "Provenance and terms", ""),
        (
            "t",
            "Source",
            "MIMIC-IV Clinical Database Demo v2.2 (PhysioNet), ODbL v1.0 - open access, "
            "no credentialing required.",
        ),
        (
            "t",
            "Attribution required",
            "Johnson A, Bulgarelli L, Pollard T, Celi LA, Mark R, Badawi O (2023). "
            "MIMIC-IV Clinical Database Demo (v2.2). PhysioNet. "
            "https://doi.org/10.13026/dp1f-ex47",
        ),
        (
            "t",
            "De-identified",
            "HIPAA Safe Harbor. Do not attempt re-identification or linkage to external "
            "identifying data.",
        ),
        ("t", "Regenerate", "python ml/evaluation/export_training_data.py"),
    ]
    r = 4
    for kind, a, b in rows:
        if kind == "h":
            ws[f"A{r}"] = a
            ws[f"A{r}"].font = Font(name=ARIAL, bold=True, size=11)
        else:
            ws[f"A{r}"] = a
            ws[f"A{r}"].font = Font(name=ARIAL, size=10)
            ws[f"B{r}"] = b
            ws[f"B{r}"].font = Font(name=ARIAL, size=10, bold=(kind == "n"))
            ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 100
    for name in ("Train", "Test", "Feature_Dictionary"):
        style(wb[name])
    wb.active = 0
    wb.save(path)
    print(f"Wrote {path}  Train={len(train):,}  Test={len(test):,}  features={len(x.columns)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
